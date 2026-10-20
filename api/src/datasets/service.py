"""Parsing, detection and row retrieval for uploaded datasets."""

from __future__ import annotations

import json
import re
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterator
from uuid import UUID

from loguru import logger

from src.config import config
from src.datasets import actions, laya_detection, storage
from src.datasets.detection import detect_column

PREVIEW_ROWS = 100
DETECTION_SAMPLE_CAP = 5000

_CURRENCY_SYMBOLS = ("$", "€", "£", "¥", "¤")
_NUMBER_CHARS_RE = re.compile(r"[0#]")
_DECIMAL_PLACES_RE = re.compile(r"\.(0+)")


def _format_number(value: float, number_format: str) -> str:
    """Render a numeric cell according to a small subset of Excel formats."""
    if not _NUMBER_CHARS_RE.search(number_format):
        return str(int(value)) if value.is_integer() else repr(value)

    symbol = ""
    for candidate in _CURRENCY_SYMBOLS:
        if candidate in number_format:
            symbol = candidate
            break

    is_percent = "%" in number_format
    has_thousands = "," in number_format

    digits_match = _DECIMAL_PLACES_RE.search(number_format)
    decimals = len(digits_match.group(1)) if digits_match else 0

    body_format = f",.{decimals}f" if has_thousands else f".{decimals}f"

    if is_percent:
        value *= 100
        return f"{value:{body_format}}%"

    return f"{symbol}{value:{body_format}}"


class DatasetNotFoundError(Exception):
    pass


def _normalize_cell(value: Any, number_format: str = "General") -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (int, float)):
        return _format_number(float(value), number_format)
    text = str(value).strip()
    return text or None


def _open_sheet_rows(
    path: Path,
) -> tuple[list[str], str, Iterator[list[Any]]]:
    """Return (sheet_names, first_sheet_name, row iterator) for a workbook."""
    from openpyxl import load_workbook

    workbook = load_workbook(str(path), read_only=True, data_only=True)
    names = list(workbook.sheetnames)
    worksheet = workbook[names[0]]

    def gen() -> Iterator[list[Any]]:
        for row in worksheet.iter_rows():
            yield [(cell.value, str(cell.number_format)) for cell in row]

    return names, names[0], gen()


def create_dataset(
    dataset_id: UUID,
    filename: str,
    suffix: str,
    content: bytes,
) -> dict[str, Any]:
    directory = storage.dataset_dir(dataset_id)
    directory.mkdir(parents=True, exist_ok=True)

    original = storage.original_path(dataset_id, suffix)
    original.write_bytes(content)

    sheet_names, sheet_name, rows_iter = _open_sheet_rows(original)
    iterator = iter(rows_iter)

    try:
        header_row = next(iterator)
    except StopIteration:
        header_row = []

    headers: list[str | None] = []
    for raw, _format in header_row:
        if raw is None:
            headers.append(None)
        else:
            text = str(raw).strip()
            headers.append(text or None)

    column_count = len(headers)
    samples: list[list[str]] = [[] for _ in range(column_count)]
    numeric_min: list[float | None] = [None] * column_count
    numeric_max: list[float | None] = [None] * column_count
    preview: list[list[str | None]] = []
    total_rows = 0

    rows_file = storage.rows_path(dataset_id)
    with rows_file.open("w", encoding="utf-8") as handle:
        for raw_row in iterator:
            row = [
                _normalize_cell(*raw_row[index])
                if index < len(raw_row)
                else None
                for index in range(column_count)
            ]
            total_rows += 1
            if len(preview) < PREVIEW_ROWS:
                preview.append(row)
            for index in range(column_count):
                value = row[index]
                if value is None:
                    continue
                if len(samples[index]) < DETECTION_SAMPLE_CAP:
                    samples[index].append(value)
                number = actions.parse_number(value)
                if number is not None:
                    low = numeric_min[index]
                    high = numeric_max[index]
                    if low is None or number < low:
                        numeric_min[index] = number
                    if high is None or number > high:
                        numeric_max[index] = number
            handle.write(json.dumps(row, ensure_ascii=False))
            handle.write("\n")

    columns: list[dict[str, Any]] = []
    for index in range(column_count):
        detected_type, sensitivity, confidence, sample_values = detect_column(
            headers[index], samples[index]
        )
        columns.append(
            {
                "index": index,
                "header": headers[index],
                "detected_type": detected_type,
                "sensitivity": sensitivity,
                "confidence": confidence,
                "sample_values": sample_values,
                "numeric_min": numeric_min[index],
                "numeric_max": numeric_max[index],
            }
        )

    # Uncertain columns go to laya-serve (heuristics stand when it is down),
    # then every column gets its sensitivity x categorizability default.
    laya_detection.enhance_columns(
        columns,
        samples,
        base_url=config.laya_base_url,
        enabled=config.laya_detect_enabled,
        min_confidence=config.laya_detect_min_confidence,
        timeout=config.laya_detect_timeout,
        api_key=config.laya_api_key,
    )
    for index, column in enumerate(columns):
        column["default_action"] = actions.column_action(
            column["detected_type"], column["sensitivity"], samples[index]
        )

    meta = {
        "id": str(dataset_id),
        "filename": filename,
        "size": len(content),
        "sheet": sheet_name,
        "sheets": sheet_names,
        "total_rows": total_rows,
        "columns": columns,
    }
    storage.meta_path(dataset_id).write_text(
        json.dumps(meta, ensure_ascii=False), encoding="utf-8"
    )
    logger.info(
        f"[DATASETS] Parsed {filename}: {column_count} columns, "
        f"{total_rows} rows, sheet '{sheet_name}'"
    )

    return {**meta, "rows": preview}


def _read_meta(dataset_id: UUID | str) -> dict[str, Any]:
    meta_file = storage.meta_path(dataset_id)
    if not meta_file.is_file():
        raise DatasetNotFoundError(str(dataset_id))
    return json.loads(meta_file.read_text(encoding="utf-8"))


def _read_window(
    dataset_id: UUID | str, offset: int, limit: int
) -> list[list[str | None]]:
    rows_file = storage.rows_path(dataset_id)
    if not rows_file.is_file():
        raise DatasetNotFoundError(str(dataset_id))

    window: list[list[str | None]] = []
    with rows_file.open("r", encoding="utf-8") as handle:
        for index, line in enumerate(handle):
            if index < offset:
                continue
            if len(window) >= limit:
                break
            window.append(json.loads(line))
    return window


def get_dataset(dataset_id: UUID | str) -> dict[str, Any]:
    meta = _read_meta(dataset_id)
    rows = _read_window(dataset_id, 0, PREVIEW_ROWS)
    return {**meta, "rows": rows}


def get_rows(
    dataset_id: UUID | str, offset: int, limit: int
) -> dict[str, Any]:
    meta = _read_meta(dataset_id)
    rows = _read_window(dataset_id, offset, limit)
    return {
        "offset": offset,
        "limit": limit,
        "total": meta["total_rows"],
        "rows": rows,
    }
