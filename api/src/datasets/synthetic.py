"""Synthetic data generation jobs backed by the SDV.

Mirrors :mod:`src.datasets.processing`: an in-memory registry tracks running
jobs, a per-job JSON file survives restarts, and the CPU-heavy SDV work runs in
a threadpool. The SDV does not expose progress callbacks, so progress advances
in coarse phases instead of per row.
"""

from __future__ import annotations

import asyncio
import json
import math
import random
import re
import threading
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any
from uuid import UUID, uuid7

from fastapi.concurrency import run_in_threadpool
from loguru import logger

from src.datasets import actions, storage
from src.datasets.dtos import SyntheticRequest
from src.datasets.service import DatasetNotFoundError

#: NRIC/FIN checksum tables (Singapore)
_NRIC_WEIGHTS = [2, 7, 6, 5, 4, 3, 2]
_NRIC_ST_TABLE = "JZIHGFEDCBA"
_NRIC_FG_TABLE = "XWUTRQPNMLK"
_NRIC_M_TABLE = "XWUTRQPNJLK"
_NRIC_RE = re.compile(r"^[STFGM]\d{7}[A-Z]$")

_GENDER_MALE = {"m", "male", "man", "boy"}
_GENDER_FEMALE = {"f", "female", "woman", "girl"}

#: Singapore address generation (match source style, source + curated SG list)
#: Curated name/type pairs. The boolean marks HDB-style blocks (``Blk``).
_SG_STREET_PAIRS = [
    ("Ang Mo Kio", "Ave", True),
    ("Bedok", "North Rd", True),
    ("Bishan", "St", True),
    ("Bukit Batok", "West Ave", True),
    ("Bukit Merah", "View", True),
    ("Bukit Panjang", "Ring Rd", True),
    ("Choa Chu Kang", "Ave", True),
    ("Clementi", "Ave", True),
    ("Geylang", "Rd", True),
    ("Hougang", "Ave", True),
    ("Jurong East", "St", True),
    ("Jurong West", "St", True),
    ("Kallang", "Bahru", True),
    ("Marine Parade", "Rd", True),
    ("Pasir Ris", "Dr", True),
    ("Punggol", "Field", True),
    ("Queenstown", "Rd", True),
    ("Sembawang", "Rd", True),
    ("Sengkang", "East Ave", True),
    ("Serangoon", "Central", True),
    ("Tampines", "Ave", True),
    ("Toa Payoh", "Lor", True),
    ("Woodlands", "Dr", True),
    ("Yishun", "Ring Rd", True),
    ("Simei", "St", True),
    ("Bugis", "St", True),
    ("Orchard", "Rd", False),
    ("Mandarin", "Drive", False),
    ("East Coast", "Rd", False),
    ("Bukit Timah", "Rd", False),
    ("Victoria", "St", False),
    ("North Bridge", "Rd", False),
    ("New Upper Changi", "Rd", False),
    ("Upper Changi", "Rd", False),
    ("West Coast", "Rd", False),
    ("Telok Blangah", "Rd", False),
]

_BLK_RE = re.compile(r"\b(?:blk|block)\s*(\d+[A-Za-z]?)\b", re.IGNORECASE)
_STREET_TYPE_RE = re.compile(
    r"\b(street|st|road|rd|avenue|ave|drive|dr|crescent|cres|lane|ln|close|cl|walk|place|terrace)\b",
    re.IGNORECASE,
)
_STREET_TYPE_MAP = {
    "st": "St",
    "street": "Street",
    "rd": "Rd",
    "road": "Road",
    "ave": "Ave",
    "avenue": "Avenue",
    "dr": "Dr",
    "drive": "Drive",
    "cres": "Cres",
    "crescent": "Crescent",
    "ln": "Ln",
    "lane": "Lane",
    "cl": "Cl",
    "close": "Close",
    "walk": "Walk",
    "place": "Place",
    "terrace": "Terrace",
}
_SG_ADDR_NORM_RE = re.compile(r"[^a-z0-9]+")

_ACTIVE: dict[str, dict[str, Any]] = {}
_LOCK = threading.Lock()

#: Detected-type -> (SDV sdtype, is-PII). The frontend lets users override
#: these suggestions before generation.
_SDTYPE_MAP: dict[str, tuple[str, bool]] = {
    "EMAIL": ("email", True),
    "PHONE": ("phone_number", True),
    "ADDRESS": ("street_address", True),
    "NAME": ("first_name", True),
    "CARD": ("credit_card_number", True),
    "NRIC": ("id", False),
    "PASSPORT": ("id", False),
    "DRIVING_LICENSE": ("id", False),
    "IP_ADDRESS": ("ipv4_address", False),
    "URL": ("categorical", False),
    "DOB": ("datetime", False),
    "DATE": ("datetime", False),
    "AGE": ("numerical", False),
    "NUMERIC": ("numerical", False),
    "TEXT": ("text", True),
    "GENDER": ("categorical", False),
    "TITLE": ("categorical", False),
    "UNKNOWN": ("categorical", False),
}

#: Sdtypes that accept SDV's ``pii`` flag. For these, ``pii=True`` means values
#: are anonymized; ``pii=False`` means the original values may be reused.
_PII_SDTYPES = {
    "address",
    "administrative_unit",
    "city",
    "country_code",
    "credit_card_number",
    "email",
    "first_name",
    "iban",
    "ipv4_address",
    "ipv6_address",
    "last_name",
    "license_plate",
    "mac_address",
    "passport_number",
    "phone_number",
    "postcode",
    "secondary_address",
    "ssn",
    "state_abbr",
    "street_address",
    "swift11",
    "swift8",
    "user_agent_string",
    "vin",
}

_DATE_FORMAT = "%Y-%m-%d"

#: Sdtypes whose generated values can be constrained to a range.
_BOUNDED_SDTYPES = {"numerical", "datetime"}

#: How many extra sampling rounds may run while looking for rows that fall
#: inside the requested bounds. The shortfall is filled with clipped rows
#: afterwards, so the requested row count is always honoured.
_MAX_RESAMPLE_ATTEMPTS = 5

#: Original values inspected when inferring display/phone formats.
_FORMAT_SAMPLE_LIMIT = 5000

#: Currency symbols the upload parser recognises in numeric cells.
_CURRENCY_SYMBOLS = ("$", "€", "£", "¥", "¤")

#: Attempts to avoid reproducing a source phone number verbatim.
_PHONE_RETRY_LIMIT = 50


@dataclass(frozen=True)
class _NumericFormat:
    """How a column's numbers were displayed in the source data."""

    prefix: str
    suffix: str
    thousands: bool
    decimals: int
    percent: bool


@dataclass(frozen=True)
class _PhoneTemplate:
    """A source phone shape with fixed digits and leading-digit choices."""

    template: str
    fixed: tuple[str | None, ...]
    leading_position: int | None
    leading: tuple[str, ...]
    weight: int


def _column_names(meta: dict[str, Any]) -> dict[int, str]:
    """Resolve unique, non-empty column names for SDV metadata."""
    names: dict[int, str] = {}
    used: set[str] = set()
    for column in meta["columns"]:
        base = (column.get("header") or "").strip() or (
            f"column_{column['index'] + 1}"
        )
        name = base
        suffix = 2
        while name in used:
            name = f"{base}_{suffix}"
            suffix += 1
        used.add(name)
        names[column["index"]] = name
    return names


def _table_name(meta: dict[str, Any]) -> str:
    name = str(meta.get("sheet") or "").strip()
    return name[:63] or "table"


def _format_bound(value: Any) -> str | None:
    """Render a detected bound as the string the form sends back."""
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _detected_bounds(column: dict[str, Any], sdtype: str) -> tuple[str | None, str | None]:
    """Observed range of a column, used to pre-fill the allowed range."""
    if sdtype == "numerical":
        return (
            _format_bound(column.get("numeric_min")),
            _format_bound(column.get("numeric_max")),
        )
    if sdtype == "datetime":
        return (
            _format_bound(column.get("datetime_min")),
            _format_bound(column.get("datetime_max")),
        )
    return None, None


def _parse_bound(text: str, sdtype: str, label: str, index: int) -> Any:
    """Parse one bound, raising ValueError with a client-facing message."""
    if sdtype == "numerical":
        number = actions.parse_number(text)
        if number is None:
            raise ValueError(
                f"Column {index}: {label} '{text}' is not a number."
            )
        return number
    try:
        return datetime.strptime(text, _DATE_FORMAT)
    except ValueError as exc:
        raise ValueError(
            f"Column {index}: {label} '{text}' is not a YYYY-MM-DD date."
        ) from exc


def resolve_bounds(request: SyntheticRequest) -> dict[int, tuple[Any, Any]]:
    """Validate the requested ranges and resolve them to comparable values.

    Returns a mapping of column index to its (min, max) pair; either side may
    be None when the user left that end open.
    """
    bounds: dict[int, tuple[Any, Any]] = {}
    for config in request.columns:
        index = config.index
        low_text = (config.min_value or "").strip()
        high_text = (config.max_value or "").strip()
        if not low_text and not high_text:
            continue
        if request.primary_key == index:
            raise ValueError(
                f"Column {index}: the primary key is generated as an id and "
                "cannot have a range."
            )
        sdtype = config.sdtype.strip()
        if sdtype not in _BOUNDED_SDTYPES:
            raise ValueError(
                f"Column {index}: ranges are only supported for numerical "
                "and datetime columns."
            )
        low = (
            _parse_bound(low_text, sdtype, "minimum", index) if low_text else None
        )
        high = (
            _parse_bound(high_text, sdtype, "maximum", index) if high_text else None
        )
        if low is not None and high is not None and low > high:
            raise ValueError(
                f"Column {index}: the minimum must not exceed the maximum."
            )
        bounds[index] = (low, high)
    return bounds


def suggest_metadata(meta: dict[str, Any]) -> dict[str, Any]:
    """Build the pre-filled metadata form for a dataset."""
    names = _column_names(meta)
    columns = []
    suggested_primary_key: int | None = None
    for column in meta["columns"]:
        sdtype, pii = _SDTYPE_MAP.get(
            column.get("detected_type", "UNKNOWN"), ("categorical", False)
        )
        if suggested_primary_key is None and sdtype == "id":
            suggested_primary_key = column["index"]
        min_value, max_value = _detected_bounds(column, sdtype)
        columns.append(
            {
                "index": column["index"],
                "header": names[column["index"]],
                "sdtype": sdtype,
                "pii": pii,
                "min_value": min_value,
                "max_value": max_value,
            }
        )
    return {
        "table_name": _table_name(meta),
        "suggested_primary_key": suggested_primary_key,
        "default_num_rows": int(meta.get("total_rows", 0)),
        "columns": columns,
    }


def start_job(
    dataset_id: UUID, meta: dict[str, Any], request: SyntheticRequest
) -> dict[str, Any]:
    job_id = str(uuid7())
    job: dict[str, Any] = {
        "job_id": job_id,
        "dataset_id": str(dataset_id),
        "status": "running",
        "progress": 0.0,
        "phase": "Starting",
        "total_rows": request.num_rows,
        "error": None,
        "export_filename": None,
        "output_rows": None,
        "request": request.model_dump(mode="json"),
    }
    with _LOCK:
        _ACTIVE[job_id] = job
    _write_status(dataset_id, job)
    asyncio.create_task(
        run_in_threadpool(_run_job, dataset_id, meta, request, job)
    )
    return dict(job)


def get_status(dataset_id: UUID | str, job_id: str) -> dict[str, Any]:
    with _LOCK:
        job = _ACTIVE.get(job_id)
        if job is not None:
            return dict(job)

    path = storage.synthetic_job_path(dataset_id, job_id)
    if not path.is_file():
        raise DatasetNotFoundError(job_id)
    data: dict[str, Any] = json.loads(path.read_text(encoding="utf-8"))
    if data.get("status") == "running":
        # The server restarted while this job was in flight.
        data["status"] = "failed"
        data["error"] = (
            "Generation was interrupted (the server restarted). Try again."
        )
        _write_status(dataset_id, data)
    return data


def export_file(
    dataset_id: UUID | str, job_id: str
) -> tuple[Path, str] | None:
    with _LOCK:
        if job_id in _ACTIVE:
            return None
    status_path = storage.synthetic_job_path(dataset_id, job_id)
    if not status_path.is_file():
        return None
    data: dict[str, Any] = json.loads(status_path.read_text(encoding="utf-8"))
    if data.get("status") != "completed":
        return None
    path = storage.synthetic_output_path(dataset_id, job_id)
    filename = data.get("export_filename") or f"{job_id}.xlsx"
    if not path.is_file():
        return None
    return path, filename


def _write_status(dataset_id: UUID | str, job: dict[str, Any]) -> None:
    path = storage.synthetic_job_path(dataset_id, job["job_id"])
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(job, ensure_ascii=False), encoding="utf-8")


def _cast_column(sdtype: str, values: list[str | None]) -> list[Any]:
    if sdtype == "numerical":
        return [
            actions.parse_number(value) if value is not None else None
            for value in values
        ]
    if sdtype == "boolean":
        lookup = {
            "true": True,
            "false": False,
            "yes": True,
            "no": False,
            "1": True,
            "0": False,
        }
        return [lookup.get((value or "").strip().lower()) for value in values]
    return values


def _display_value(value: Any) -> str | None:
    """Render a generated cell the same way uploaded dataset rows are stored."""
    import pandas as pd

    if value is None:
        return None
    try:
        if pd.isna(value):
            return None
    except (TypeError, ValueError):
        pass
    if hasattr(value, "item") and not isinstance(value, (str, bytes)):
        value = value.item()
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).strip()
    return text or None


def _numeric_format(values: list[str | None]) -> _NumericFormat | None:
    """Infer the display shape of a numeric column from its source strings."""
    shapes: dict[tuple[str, str, bool, bool], list[int]] = {}
    checked = 0
    for raw in values:
        if checked >= _FORMAT_SAMPLE_LIMIT:
            break
        if raw is None:
            continue
        text = str(raw).strip()
        if not text or actions.parse_number(text) is None:
            continue

        percent = text.endswith("%")
        body = text[:-1].strip() if percent else text
        prefix = ""
        for symbol in _CURRENCY_SYMBOLS:
            if body.startswith(symbol):
                prefix = symbol
                body = body[len(symbol) :].strip()
                break

        thousands = "," in body
        decimal_match = re.search(r"\.(\d+)", body)
        decimals = len(decimal_match.group(1)) if decimal_match else 0

        key = (prefix, "%" if percent else "", thousands, percent)
        entry = shapes.setdefault(key, [0, 0])
        entry[0] += 1
        entry[1] = max(entry[1], decimals)
        checked += 1

    if not shapes:
        return None

    (prefix, suffix, thousands, percent), (_, decimals) = max(
        shapes.items(), key=lambda item: item[1][0]
    )
    return _NumericFormat(
        prefix=prefix,
        suffix=suffix,
        thousands=thousands,
        decimals=decimals,
        percent=percent,
    )


def _format_numeric(value: Any, number_format: _NumericFormat) -> str | None:
    """Render a generated number the way the source column was displayed."""
    import pandas as pd

    try:
        if pd.isna(value):
            return None
    except (TypeError, ValueError):
        pass

    try:
        number = float(value)
    except (TypeError, ValueError):
        text = str(value).strip()
        return text or None

    sign = "-" if number < 0 else ""
    body_format = (
        f",.{number_format.decimals}f"
        if number_format.thousands
        else f".{number_format.decimals}f"
    )
    body = f"{abs(number):{body_format}}"
    return f"{sign}{number_format.prefix}{body}{number_format.suffix}"


def _excel_number_format(number_format: _NumericFormat | None) -> str | None:
    """Translate an inferred display shape into an Excel number format."""
    if number_format is None:
        return None
    decimals = "." + ("0" * number_format.decimals) if number_format.decimals else ""
    body = ("#,##0" if number_format.thousands else "0") + decimals
    if number_format.percent:
        return f"{body}%"
    if number_format.prefix:
        return f'"{number_format.prefix}"{body}'
    return body


def _write_excel(
    frame: Any,
    path: Path,
    names: dict[int, str],
    request: SyntheticRequest,
    formats: dict[str, _NumericFormat],
) -> None:
    """Write the export workbook, keeping numeric cells but restoring formats."""
    import pandas as pd
    from openpyxl.utils import get_column_letter

    ordered = frame[[names[config.index] for config in request.columns]].copy()
    for name, number_format in formats.items():
        if number_format.percent and name in ordered.columns:
            ordered[name] = pd.to_numeric(ordered[name], errors="coerce") / 100

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        ordered.to_excel(writer, index=False, sheet_name="Sheet1")
        worksheet = writer.sheets["Sheet1"]
        for offset, config in enumerate(request.columns, start=1):
            number_format = _excel_number_format(
                formats.get(names[config.index])
            )
            if number_format is None:
                continue
            letter = get_column_letter(offset)
            worksheet.column_dimensions[letter].number_format = number_format
            for row in range(2, len(ordered) + 2):
                worksheet.cell(row=row, column=offset).number_format = number_format


def _write_rows(frame: Any, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as handle:
        for row in frame.itertuples(index=False, name=None):
            handle.write(
                json.dumps(
                    [_display_value(value) for value in row],
                    ensure_ascii=False,
                )
            )
            handle.write("\n")


def _load_dataframe(
    dataset_id: UUID, meta: dict[str, Any], request: SyntheticRequest
) -> tuple[Any, dict[int, list[str | None]]]:
    """Load all dataset rows and cast them according to the chosen sdtypes."""
    import pandas as pd

    configs = {column.index: column for column in request.columns}
    indexes = [column.index for column in request.columns]
    names = _column_names(meta)
    column_meta = {column["index"]: column for column in meta["columns"]}

    raw: dict[int, list[str | None]] = {index: [] for index in indexes}
    with storage.rows_path(dataset_id).open("r", encoding="utf-8") as handle:
        for line in handle:
            row = json.loads(line)
            for index in indexes:
                raw[index].append(row[index] if index < len(row) else None)

    data: dict[str, Any] = {}
    for index, config in configs.items():
        if index not in column_meta:
            raise ValueError(f"Unknown column index: {index}")
        data[names[index]] = _cast_column(config.sdtype, raw[index])

    return pd.DataFrame(data), raw


def _build_metadata(meta: dict[str, Any], request: SyntheticRequest) -> Any:
    from sdv.metadata import Metadata

    names = _column_names(meta)
    metadata = Metadata()
    metadata.add_table(table_name=request.table_name)

    for config in request.columns:
        name = names[config.index]
        sdtype = config.sdtype.strip()
        # A primary key must use an id-like sdtype, so override the user's
        # choice for that column (it is generated as fresh identifiers).
        if request.primary_key == config.index:
            sdtype = "id"
            pii = False
        else:
            pii = config.pii
        kwargs: dict[str, Any] = {}
        if sdtype == "datetime":
            kwargs["datetime_format"] = _DATE_FORMAT
        if sdtype in _PII_SDTYPES:
            kwargs["pii"] = pii
        if sdtype == "text":
            # SDV defaults text to pii=True (generate new sentences); passing
            # pii=False would reuse the original values verbatim, so never let
            # the form disable it.
            kwargs["pii"] = True
        metadata.add_column(
            column_name=name, sdtype=sdtype, table_name=request.table_name, **kwargs
        )

    if request.primary_key is not None:
        metadata.set_primary_key(
            column_name=names[request.primary_key], table_name=request.table_name
        )

    metadata.validate()
    return metadata


def _integer_columns(data: Any) -> set[str]:
    """Names of the numerical columns whose source values are whole numbers."""
    import pandas as pd

    names: set[str] = set()
    for name in data.columns:
        values = pd.to_numeric(data[name], errors="coerce").dropna()
        if len(values) and bool((values % 1 == 0).all()):
            names.add(name)
    return names


def _effective_bounds(
    bounds: dict[int, tuple[Any, Any]], names: dict[int, str], integers: set[str]
) -> dict[int, tuple[Any, Any]]:
    """Tighten the bounds of whole-number columns to whole numbers."""
    effective: dict[int, tuple[Any, Any]] = {}
    for index, (low, high) in bounds.items():
        is_date = isinstance(low, datetime) or isinstance(high, datetime)
        if names[index] in integers and not is_date:
            new_low = math.ceil(low) if low is not None else None
            new_high = math.floor(high) if high is not None else None
            if (
                new_low is None
                or new_high is None
                or new_low <= new_high
            ):
                low, high = new_low, new_high
        effective[index] = (low, high)
    return effective


def _to_datetime_series(series: Any) -> Any:
    import pandas as pd

    if pd.api.types.is_datetime64_any_dtype(series):
        return series
    values = pd.to_datetime(series, format=_DATE_FORMAT, errors="coerce")
    if int(values.notna().sum()) < int(series.notna().sum()):
        # Not every value is an ISO string (SDV may hand back real datetimes),
        # so fall back to the general parser when it recovers more values.
        loose = pd.to_datetime(series, errors="coerce")
        if int(loose.notna().sum()) > int(values.notna().sum()):
            return loose
    return values


def _bounds_mask(
    frame: Any,
    bounds: dict[int, tuple[Any, Any]],
    names: dict[int, str],
) -> Any:
    """Boolean mask of the rows whose values fall inside every range."""
    import pandas as pd

    mask = pd.Series(True, index=frame.index)
    for index, (low, high) in bounds.items():
        name = names[index]
        if name not in frame.columns:
            continue
        series = frame[name]
        if isinstance(low, datetime) or isinstance(high, datetime):
            values = _to_datetime_series(series)
        else:
            values = pd.to_numeric(series, errors="coerce")
        within = pd.Series(True, index=frame.index)
        if low is not None:
            within &= values.isna() | (values >= low)
        if high is not None:
            within &= values.isna() | (values <= high)
        mask &= within
    return mask


def _clip_frame(
    frame: Any,
    bounds: dict[int, tuple[Any, Any]],
    names: dict[int, str],
    integers: set[str],
) -> Any:
    """Force every value of the bounded columns inside its range."""
    import pandas as pd

    clipped = frame.copy()
    for index, (low, high) in bounds.items():
        name = names[index]
        if name not in clipped.columns:
            continue
        series = clipped[name]
        if isinstance(low, datetime) or isinstance(high, datetime):
            values = _to_datetime_series(series)
            lower = pd.Timestamp(low) if low is not None else None
            upper = pd.Timestamp(high) if high is not None else None
            values = values.clip(lower=lower, upper=upper)
            if pd.api.types.is_datetime64_any_dtype(series):
                clipped[name] = values
            else:
                clipped[name] = values.dt.strftime(_DATE_FORMAT)
            continue

        values = pd.to_numeric(series, errors="coerce")
        if name in integers:
            values = values.round()
        values = values.clip(lower=low, upper=high)
        if name in integers:
            if bool(values.isna().any()):
                clipped[name] = values.astype("float64")
            else:
                clipped[name] = values.astype("int64")
        else:
            clipped[name] = values
    return clipped


def _nric_checksum(prefix: str, digits7: str) -> str:
    """Return the checksum letter for a 7-digit NRIC/FIN body."""
    prefix = prefix.upper()
    total = sum(int(d) * w for d, w in zip(digits7, _NRIC_WEIGHTS))
    if prefix == "M":
        return _NRIC_M_TABLE[(total + 3) % 11]
    if prefix in {"T", "G"}:
        total += 4
    table = _NRIC_FG_TABLE if prefix in {"F", "G"} else _NRIC_ST_TABLE
    return table[total % 11]


def _generate_nric(prefix: str, want_male: bool | None, existing: set[str]) -> str:
    """Generate a checksum-valid, unique NRIC/FIN."""
    for _ in range(10000):
        digits = [str(random.randint(0, 9)) for _ in range(6)]
        if want_male is None:
            digits.append(str(random.randint(0, 9)))
        else:
            choices = [1, 3, 5, 7, 9] if want_male else [0, 2, 4, 6, 8]
            digits.append(str(random.choice(choices)))
        digits7 = "".join(digits)
        nric = f"{prefix}{digits7}{_nric_checksum(prefix, digits7)}"
        if nric not in existing:
            return nric
    raise ValueError("Unable to generate a unique NRIC after 10000 attempts")


def _normalize_gender(value: Any) -> bool | None:
    """Map a generated gender value to ``True`` for male, ``False`` for female."""
    if value is None:
        return None
    token = re.sub(r"[^a-z]", "", str(value).lower())
    if token in _GENDER_MALE:
        return True
    if token in _GENDER_FEMALE:
        return False
    return None


def _gender_parities(
    frame: Any, meta: dict[str, Any], names: dict[int, str]
) -> list[bool | None]:
    """Return the per-row gender parity, if a gender column was generated."""
    for column in meta.get("columns", []):
        if column.get("detected_type") != "GENDER":
            continue
        index = column["index"]
        name = names.get(index)
        if name and name in frame.columns:
            return [_normalize_gender(value) for value in frame[name]]
    return [None] * len(frame)


def _nric_prefix_pool(original_df: Any, name: str) -> dict[str, int]:
    """Count the source prefixes so generated IDs mirror the original mix."""
    if original_df is None or name not in getattr(original_df, "columns", []):
        return {"S": 1, "T": 1}
    counts: dict[str, int] = {}
    for value in original_df[name].dropna().astype(str):
        normalized = value.strip().upper().replace(" ", "")
        if _NRIC_RE.match(normalized):
            counts[normalized[0]] = counts.get(normalized[0], 0) + 1
    return counts or {"S": 1, "T": 1}


def _normalize_sg_addr(value: str) -> str:
    return _SG_ADDR_NORM_RE.sub("", value.lower())


def _parse_sg_address(value: Any) -> tuple[str, str, bool] | None:
    """Extract (street name, street type, is-HDB) from a source address."""
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None

    block_match = _BLK_RE.search(text)
    type_match = _STREET_TYPE_RE.search(text)
    if type_match is None:
        return None

    street_type = _STREET_TYPE_MAP.get(type_match.group(1).lower())
    if street_type is None:
        return None

    start = 0
    if block_match:
        start = block_match.end()
    else:
        house_match = re.match(r"^\s*\d+[A-Za-z]?\b", text)
        if house_match:
            start = house_match.end()

    street_name = " ".join(text[start : type_match.start()].strip(" ,.-").split())
    if not street_name:
        return None

    return street_name, street_type, bool(block_match)


def _sg_address_pools(original_df: Any, name: str) -> dict[str, Any]:
    source_pairs: list[tuple[str, str, bool]] = []
    if original_df is not None and name in getattr(original_df, "columns", []):
        for value in original_df[name].dropna().astype(str):
            parsed = _parse_sg_address(value)
            if parsed:
                source_pairs.append(parsed)
    return {"source_pairs": source_pairs, "curated_pairs": list(_SG_STREET_PAIRS)}


def _compose_sg_address(street_name: str, street_type: str, use_blk: bool) -> str:
    if use_blk:
        block_num = random.randint(1, 999)
        suffix = random.choice(["", "", "", "A", "B", "C"])
        block = f"{block_num}{suffix}"
        street_num = random.randint(1, 200)
        return f"Blk {block} {street_name} {street_type} {street_num}"
    house_num = random.randint(1, 999)
    suffix = random.choice(["", "", "", "A", "B"])
    house = f"{house_num}{suffix}"
    return f"{house} {street_name} {street_type}"


def _generate_sg_address(pools: dict[str, Any], used: set[str]) -> str:
    source_pairs = pools.get("source_pairs") or []
    curated_pairs = pools.get("curated_pairs") or []

    for _ in range(20000):
        if source_pairs and (not curated_pairs or random.random() < 0.7):
            street_name, street_type, use_blk = random.choice(source_pairs)
        else:
            street_name, street_type, use_blk = random.choice(curated_pairs)

        addr = _compose_sg_address(street_name, street_type, use_blk)
        norm = _normalize_sg_addr(addr)
        if norm not in used:
            used.add(norm)
            return addr

    # Fallback: add a unit number so a unique value can still be produced.
    pairs = source_pairs or curated_pairs
    while True:
        street_name, street_type, _ = random.choice(pairs)
        addr = (
            f"{_compose_sg_address(street_name, street_type, False)} "
            f"#01-{random.randint(1, 999):03d}"
        )
        norm = _normalize_sg_addr(addr)
        if norm not in used:
            used.add(norm)
            return addr


def _postprocess_address(
    synthetic: Any,
    original_df: Any,
    meta: dict[str, Any],
    request: SyntheticRequest,
    names: dict[int, str],
) -> Any:
    """Replace ADDRESS columns with Singapore-style addresses."""
    import pandas as pd

    if synthetic is None or len(synthetic) == 0:
        return synthetic
    frame = synthetic.copy()
    column_meta = {column["index"]: column for column in meta.get("columns", [])}

    for config in request.columns:
        index = config.index
        if column_meta.get(index, {}).get("detected_type") != "ADDRESS":
            continue
        if request.primary_key == index:
            continue
        if config.sdtype.strip() != "street_address" or not config.pii:
            continue
        name = names.get(index)
        if name is None or name not in frame.columns:
            continue

        pools = _sg_address_pools(original_df, name)
        used: set[str] = set()
        if original_df is not None and name in getattr(original_df, "columns", []):
            used.update(
                _normalize_sg_addr(str(value))
                for value in original_df[name].dropna()
            )
        generated: list[Any] = []
        for value in frame[name]:
            if pd.isna(value):
                generated.append(value)
                continue
            generated.append(_generate_sg_address(pools, used))
        frame[name] = generated

    return frame


def _phone_digits(value: Any) -> str:
    return re.sub(r"\D", "", str(value))


def _phone_templates(values: Any) -> list[_PhoneTemplate]:
    """Infer the source phone shapes, preserving fixed digits and separators."""
    grouped: dict[str, list[str]] = {}
    checked = 0
    for raw in values:
        if checked >= _FORMAT_SAMPLE_LIMIT:
            break
        if raw is None:
            continue
        text = str(raw).strip()
        if not text:
            continue
        digits = _phone_digits(text)
        if not 7 <= len(digits) <= 15:
            continue
        template = re.sub(r"\d", "9", text)
        grouped.setdefault(template, []).append(digits)
        checked += 1

    profiles: list[_PhoneTemplate] = []
    for template, digit_values in grouped.items():
        length = len(digit_values[0])
        if any(len(value) != length for value in digit_values):
            continue
        columns = [set(value[index] for value in digit_values) for index in range(length)]
        fixed = tuple(
            next(iter(choices)) if len(choices) == 1 else None for choices in columns
        )
        variable_positions = [
            index for index, choices in enumerate(columns) if len(choices) > 1
        ]
        leading_position = variable_positions[0] if variable_positions else None
        leading = (
            tuple(sorted(value[leading_position] for value in digit_values))
            if leading_position is not None
            else ()
        )
        profiles.append(
            _PhoneTemplate(
                template=template,
                fixed=fixed,
                leading_position=leading_position,
                leading=leading,
                weight=len(digit_values),
            )
        )
    return profiles


def _render_phone_digits(profile: _PhoneTemplate, digits: list[str]) -> str:
    rendered: list[str] = []
    digit_index = 0
    for char in profile.template:
        if char == "9":
            rendered.append(digits[digit_index])
            digit_index += 1
        else:
            rendered.append(char)
    return "".join(rendered)


def _render_phone(profile: _PhoneTemplate) -> tuple[str, str]:
    digits: list[str] = []
    for position, fixed in enumerate(profile.fixed):
        if fixed is not None:
            digits.append(fixed)
        elif position == profile.leading_position and profile.leading:
            digits.append(random.choice(profile.leading))
        else:
            digits.append(str(random.randint(0, 9)))
    return _render_phone_digits(profile, digits), "".join(digits)


def _generate_phone(profiles: list[_PhoneTemplate], forbidden: set[str]) -> str:
    """Generate a phone number matching the source shape, avoiding exact reuse."""
    profile = profiles[0]
    rendered = ""
    digits = ""
    for _ in range(_PHONE_RETRY_LIMIT):
        profile = random.choices(
            profiles, weights=[profile.weight for profile in profiles], k=1
        )[0]
        rendered, digits = _render_phone(profile)
        if digits not in forbidden:
            return rendered

    if not digits or digits not in forbidden:
        return rendered

    mutable = list(digits)
    position = next(
        (
            index
            for index in range(len(mutable) - 1, -1, -1)
            if profile.fixed[index] is None
        ),
        len(mutable) - 1,
    )
    for _ in range(10):
        mutable[position] = str(random.randint(0, 9))
        candidate = "".join(mutable)
        if candidate not in forbidden:
            return _render_phone_digits(profile, mutable)
    return _render_phone_digits(profile, mutable)


def _postprocess_phone(
    synthetic: Any,
    original_df: Any,
    meta: dict[str, Any],
    request: SyntheticRequest,
    names: dict[int, str],
) -> Any:
    """Replace phone numbers with values that follow the source format."""
    import pandas as pd

    if synthetic is None or len(synthetic) == 0:
        return synthetic
    frame = synthetic.copy()
    column_meta = {column["index"]: column for column in meta.get("columns", [])}

    for config in request.columns:
        index = config.index
        if column_meta.get(index, {}).get("detected_type") != "PHONE":
            continue
        if request.primary_key == index:
            continue
        if config.sdtype.strip() != "phone_number" or not config.pii:
            continue
        name = names.get(index)
        if name is None or name not in frame.columns:
            continue
        if original_df is None or name not in getattr(original_df, "columns", []):
            continue

        source = original_df[name]
        profiles = _phone_templates(source)
        if not profiles:
            continue
        forbidden = {
            _phone_digits(value)
            for value in source.dropna()
            if 7 <= len(_phone_digits(value)) <= 15
        }

        generated: list[Any] = []
        for value in frame[name]:
            if pd.isna(value):
                generated.append(value)
                continue
            generated.append(_generate_phone(profiles, forbidden))
        frame[name] = generated

    return frame


def _postprocess_nric(
    synthetic: Any,
    original_df: Any,
    meta: dict[str, Any],
    request: SyntheticRequest,
    names: dict[int, str],
) -> Any:
    """Replace identifier-style NRIC columns with valid, realistic values."""
    if synthetic is None or len(synthetic) == 0:
        return synthetic
    frame = synthetic.copy()
    parities = _gender_parities(frame, meta, names)
    column_meta = {column["index"]: column for column in meta.get("columns", [])}

    for config in request.columns:
        index = config.index
        if column_meta.get(index, {}).get("detected_type") != "NRIC":
            continue
        is_primary = request.primary_key == index
        if not (is_primary or config.sdtype.strip() == "id"):
            continue
        name = names.get(index)
        if name is None or name not in frame.columns:
            continue

        prefix_counts = _nric_prefix_pool(original_df, name)
        prefixes = list(prefix_counts)
        weights = [prefix_counts[prefix] for prefix in prefixes]
        used: set[str] = set()
        generated: list[str] = []
        for parity in parities:
            prefix = random.choices(prefixes, weights=weights, k=1)[0]
            nric = _generate_nric(prefix, parity, used)
            used.add(nric)
            generated.append(nric)
        frame[name] = generated

    return frame


def _pin_uniform_marginals(synthesizer: Any, columns: set[str]) -> None:
    """Pin categorical marginals to [0, 1] so source frequencies are preserved."""
    model = getattr(synthesizer, "_model", None)
    univariates = getattr(model, "univariates", None)
    if model is None or univariates is None:
        return
    for index, name in enumerate(getattr(model, "columns", [])):
        if name in columns and index < len(univariates):
            univariates[index]._params = {"loc": 0.0, "scale": 1.0}


def _sample_within_bounds(
    synthesizer: Any,
    request: SyntheticRequest,
    bounds: dict[int, tuple[Any, Any]],
    names: dict[int, str],
    integers: set[str],
    job: dict[str, Any],
) -> Any:
    """Sample rows, preferring values that respect the requested ranges.

    Rows outside the ranges are dropped and sampled again; if the budget runs
    out the shortfall is filled with clipped rows so the requested row count
    is always met.
    """
    import pandas as pd

    target = request.num_rows
    if not bounds:
        return synthesizer.sample(num_rows=target)[request.table_name]

    kept: list[Any] = []
    collected = 0
    for _ in range(_MAX_RESAMPLE_ATTEMPTS):
        if collected >= target:
            break
        batch = synthesizer.sample(num_rows=target - collected)[
            request.table_name
        ]
        valid = batch[_bounds_mask(batch, bounds, names)]
        if len(valid):
            kept.append(valid)
            collected += len(valid)
        with _LOCK:
            job.update(phase="Enforcing bounds", progress=0.8)

    if collected < target:
        filler = synthesizer.sample(num_rows=target - collected)[
            request.table_name
        ]
        kept.append(filler)

    frame = pd.concat(kept, ignore_index=True) if len(kept) > 1 else kept[0]
    frame = frame.head(target).reset_index(drop=True)
    return _clip_frame(frame, bounds, names, integers)


def _run_job(
    dataset_id: UUID,
    meta: dict[str, Any],
    request: SyntheticRequest,
    job: dict[str, Any],
) -> None:
    job_id = job["job_id"]
    try:
        with _LOCK:
            job.update(phase="Loading data", progress=0.05)
        data, raw = _load_dataframe(dataset_id, meta, request)
        names = _column_names(meta)
        numeric_formats: dict[str, _NumericFormat] = {}
        for cfg in request.columns:
            if request.primary_key == cfg.index:
                continue
            if cfg.sdtype.strip() != "numerical":
                continue
            cname = names.get(cfg.index)
            if not cname:
                continue
            number_format = _numeric_format(raw.get(cfg.index, []))
            if number_format is not None:
                numeric_formats[cname] = number_format
        del raw

        with _LOCK:
            job.update(phase="Preparing metadata", progress=0.25)
        metadata = _build_metadata(meta, request)

        with _LOCK:
            job.update(phase="Training model", progress=0.4)
        from sdv.single_table import GaussianCopulaSynthesizer

        numerical_distributions: dict[str, str] = {}
        for cfg in request.columns:
            if request.primary_key == cfg.index:
                continue
            sdtype = cfg.sdtype.strip()
            if sdtype in ("categorical", "boolean"):
                cname = names.get(cfg.index)
                if cname:
                    numerical_distributions[cname] = "uniform"
        synthesizer = GaussianCopulaSynthesizer(
            metadata, numerical_distributions=numerical_distributions or None
        )
        synthesizer.fit({request.table_name: data})
        _pin_uniform_marginals(synthesizer, set(numerical_distributions))
        if hasattr(synthesizer, "_set_random_state"):
            synthesizer._set_random_state(random.randint(0, 2**31 - 1))

        with _LOCK:
            job.update(phase="Generating synthetic data", progress=0.75)
        integers = _integer_columns(data)
        bounds = _effective_bounds(resolve_bounds(request), names, integers)
        synthetic = _sample_within_bounds(
            synthesizer, request, bounds, names, integers, job
        )
        synthetic = _postprocess_nric(synthetic, data, meta, request, names)
        synthetic = _postprocess_address(synthetic, data, meta, request, names)
        synthetic = _postprocess_phone(synthetic, data, meta, request, names)

        with _LOCK:
            job.update(phase="Writing file", progress=0.9)
        export_path = storage.synthetic_output_path(dataset_id, job_id)
        export_path.parent.mkdir(parents=True, exist_ok=True)
        _write_excel(synthetic, export_path, names, request, numeric_formats)

        ordered = synthetic[
            [names[config.index] for config in request.columns]
        ].copy()
        for cname, number_format in numeric_formats.items():
            if cname in ordered.columns:
                ordered[cname] = [
                    _format_numeric(value, number_format)
                    for value in ordered[cname]
                ]
        _write_rows(ordered, storage.synthetic_rows_path(dataset_id, job_id))

        stem = Path(str(meta.get("filename") or "dataset")).stem
        export_filename = f"{stem}_synthetic.xlsx"
        with _LOCK:
            job.update(
                status="completed",
                phase="Done",
                progress=1.0,
                output_rows=len(synthetic),
                export_filename=export_filename,
            )
        logger.info(
            f"[SYNTHETIC] job {job_id}: generated {len(synthetic)} rows"
        )
    except Exception as exc:  # noqa: BLE001 - surface any failure to the client
        logger.exception(f"[SYNTHETIC] job {job_id} failed")
        with _LOCK:
            job.update(status="failed", phase="Failed", error=str(exc))
    finally:
        with _LOCK:
            snapshot = dict(job)
        _write_status(dataset_id, snapshot)
        with _LOCK:
            _ACTIVE.pop(job_id, None)
